#!/usr/bin/env python3
"""بررسی سلامت چیدمان فایل pptx بدون نیاز به PowerPoint.

کارهایی که انجام می‌دهد:
  * باز کردن فایل و اطمینان از خوانا بودن آن
  * بیرون‌زدگی شکل‌های متنی از محدوده‌ی اسلاید
  * هم‌پوشانی کادرهای متنی (که معمولاً به معنی متن روی متن است)

اجرا:
    python3 validate_deck.py output/deck.pptx
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

TOLERANCE = Emu(int(0.02 * 914400))  # ۲ صدم اینچ


def rect(shape):
    return (shape.left, shape.top, shape.left + shape.width, shape.top + shape.height)


def overlap(a, b) -> bool:
    """هم‌پوشانی واقعی (نه فقط چسبیدن لبه‌ها)."""
    return not (
        a[2] <= b[0] + TOLERANCE
        or b[2] <= a[0] + TOLERANCE
        or a[3] <= b[1] + TOLERANCE
        or b[3] <= a[1] + TOLERANCE
    )


def has_text(shape) -> bool:
    return shape.has_text_frame and bool(shape.text_frame.text.strip())


def validate(path: Path) -> int:
    prs = Presentation(path)
    sw, sh = prs.slide_width, prs.slide_height
    problems = 0

    print(f"بررسی {path}  ({len(prs.slides)} اسلاید)")
    for idx, slide in enumerate(prs.slides, start=1):
        shapes = list(slide.shapes)
        text_shapes = [s for s in shapes if has_text(s)]

        for shape in shapes:
            left, top, right, bottom = rect(shape)
            if has_text(shape):
                if left < -TOLERANCE or top < -TOLERANCE:
                    problems += 1
                    print(f"  [!] اسلاید {idx}: کادر متنی از لبه بیرون زده -> {shape.name}")
                if right > sw + TOLERANCE or bottom > sh + TOLERANCE:
                    problems += 1
                    print(f"  [!] اسلاید {idx}: کادر متنی از لبه بیرون زده -> {shape.name}")

        for i, a in enumerate(text_shapes):
            for b in text_shapes[i + 1:]:
                if overlap(rect(a), rect(b)):
                    problems += 1
                    ta = a.text_frame.text.strip().splitlines()[0][:28]
                    tb = b.text_frame.text.strip().splitlines()[0][:28]
                    print(f"  [!] اسلاید {idx}: هم‌پوشانی متن «{ta}» با «{tb}»")

    if problems:
        print(f"\n{problems} هشدار پیدا شد (باید دستی بررسی شوند).")
        return 1
    print("\nهمه‌چیز مرتب است: بدون بیرون‌زدگی و بدون هم‌پوشانی متن.")
    return 0


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "output/deck.pptx")
    sys.exit(validate(target))
